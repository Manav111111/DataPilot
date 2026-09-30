import csv
import io
import json
import os
import re
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.core.config import settings


class DatasetExportService:
    """
    Robust export engine generating CSV, Excel (.xlsx), and JSON exports
    with CSV formula injection defense and styled multi-sheet Excel reports.
    """

    @classmethod
    def sanitize_csv_cell(cls, val: Any) -> str:
        """
        Mitigates CSV formula injection attacks by prepending a single quote
        to cells starting with =, +, -, @, tab, or carriage return.
        """
        if val is None:
            return ""
        s = str(val)
        if s.startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + s
        return s

    @classmethod
    def ensure_storage_dir(cls) -> Path:
        p = Path(settings.EXPORT_STORAGE_DIR)
        p.mkdir(parents=True, exist_ok=True)
        return p

    @classmethod
    def cleanup_old_exports(cls):
        """
        Removes export files older than retention hours.
        """
        try:
            storage_dir = cls.ensure_storage_dir()
            cutoff = time.time() - (settings.EXPORT_FILE_RETENTION_HOURS * 3600)
            for f in storage_dir.glob("export_*"):
                if f.is_file() and f.stat().st_mtime < cutoff:
                    try:
                        f.unlink()
                    except Exception:
                        pass
        except Exception:
            pass

    @classmethod
    def generate_csv(
        cls,
        records: List[Dict[str, Any]],
        columns: Optional[List[str]] = None,
        include_provenance: bool = True,
        include_warnings: bool = False,
    ) -> str:
        """
        Generates CSV string with formula injection protection.
        """
        if not records:
            return ""

        # Determine columns
        if not columns:
            all_cols = set()
            for r in records:
                data = r.get("record_data") or r.get("data") or {}
                all_cols.update(data.keys())
            columns = sorted(list(all_cols))

        header = list(columns)
        if include_warnings:
            header.append("validation_status")
        if include_provenance:
            header.append("source_count")

        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(header)

        for r in records:
            data = r.get("record_data") or r.get("data") or {}
            row = [cls.sanitize_csv_cell(data.get(c, "")) for c in columns]
            if include_warnings:
                row.append(r.get("validation_status", "valid"))
            if include_provenance:
                row.append(str(r.get("source_count", 1)))
            writer.writerow(row)

        return output.getvalue()

    @classmethod
    def generate_excel(
        cls,
        records: List[Dict[str, Any]],
        dataset_name: str = "Dataset",
        columns: Optional[List[str]] = None,
        quality_summary: Optional[Dict[str, Any]] = None,
        include_provenance: bool = True,
        include_warnings: bool = False,
    ) -> bytes:
        """
        Generates formatted .xlsx Excel workbook with frozen header row,
        custom styling, and an optional quality overview sheet.
        """
        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = "Data Records"

        # Determine columns
        if not columns:
            all_cols = set()
            for r in records:
                data = r.get("record_data") or r.get("data") or {}
                all_cols.update(data.keys())
            columns = sorted(list(all_cols))

        headers = [c.replace("_", " ").title() for c in columns]
        if include_warnings:
            headers.append("Validation Status")
        if include_provenance:
            headers.append("Source Count")

        # Header styling
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")  # Indigo-900
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws_data.append(headers)
        ws_data.row_dimensions[1].height = 28

        for col_num in range(1, len(headers) + 1):
            cell = ws_data.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Add data rows
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        row_font = Font(name="Calibri", size=10)

        for row_idx, r in enumerate(records, start=2):
            data = r.get("record_data") or r.get("data") or {}
            row_vals = [data.get(c, "") for c in columns]
            if include_warnings:
                row_vals.append(r.get("validation_status", "valid"))
            if include_provenance:
                row_vals.append(r.get("source_count", 1))

            ws_data.append(row_vals)
            ws_data.row_dimensions[row_idx].height = 20

            # Alternate row striping
            row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if row_idx % 2 == 0 else PatternFill(fill_type=None)
            for col_idx in range(1, len(row_vals) + 1):
                c_cell = ws_data.cell(row=row_idx, column=col_idx)
                c_cell.font = row_font
                c_cell.border = thin_border
                if row_idx % 2 == 0:
                    c_cell.fill = row_fill

        # Freeze top row
        ws_data.freeze_panes = "A2"

        # Auto-adjust column widths
        for col in ws_data.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = min(len(val_str), 50)
            ws_data.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Quality & Metadata Sheet
        if quality_summary:
            ws_meta = wb.create_sheet(title="Quality Report")
            ws_meta.append(["DataNexa Dataset Quality & Metadata Summary"])
            ws_meta.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="312E81")
            ws_meta.append([])
            ws_meta.append(["Dataset Name", dataset_name])
            ws_meta.append(["Exported At", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")])
            ws_meta.append(["Total Records", len(records)])
            ws_meta.append(["Overall Quality Score", f"{quality_summary.get('overall_score', 0)} / 100"])

            dim_scores = quality_summary.get("dimension_scores", {})
            for dim, score in dim_scores.items():
                ws_meta.append([f"Quality: {dim.capitalize()}", f"{score} / 100"])

            for r_idx in range(3, 10):
                ws_meta.cell(row=r_idx, column=1).font = Font(bold=True)
            ws_meta.column_dimensions["A"].width = 28
            ws_meta.column_dimensions["B"].width = 35

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def generate_json(
        cls,
        records: List[Dict[str, Any]],
        columns: Optional[List[str]] = None,
        include_provenance: bool = True,
        as_jsonl: bool = False,
    ) -> str:
        """
        Generates clean JSON string or JSON Lines.
        """
        formatted_list = []
        for r in records:
            data = r.get("record_data") or r.get("data") or {}
            if columns:
                row_dict = {c: data.get(c) for c in columns}
            else:
                row_dict = dict(data)

            if include_provenance:
                row_dict["_provenance"] = {
                    "source_count": r.get("source_count", 1),
                    "validation_status": r.get("validation_status", "valid"),
                    "record_hash": r.get("record_hash"),
                }
            formatted_list.append(row_dict)

        if as_jsonl:
            return "\n".join(json.dumps(row, default=str) for row in formatted_list)
        return json.dumps(formatted_list, indent=2, default=str)

    @classmethod
    def export_to_file(
        cls,
        records: List[Dict[str, Any]],
        export_format: str,
        dataset_id: str,
        dataset_name: str = "Dataset",
        columns: Optional[List[str]] = None,
        quality_summary: Optional[Dict[str, Any]] = None,
        include_provenance: bool = True,
        include_warnings: bool = False,
    ) -> Dict[str, Any]:
        """
        Exports records to a secure local file and returns file reference and size.
        """
        cls.cleanup_old_exports()
        storage_dir = cls.ensure_storage_dir()
        export_id = str(uuid.uuid4())
        ext = export_format.lower()
        if ext == "excel":
            ext = "xlsx"
        filename = f"export_{dataset_id[:8]}_{export_id[:8]}.{ext}"
        filepath = storage_dir / filename

        if ext == "csv":
            content = cls.generate_csv(records, columns, include_provenance, include_warnings)
            with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
                f.write(content)
        elif ext == "xlsx":
            bytes_content = cls.generate_excel(records, dataset_name, columns, quality_summary, include_provenance, include_warnings)
            with open(filepath, "wb") as f:
                f.write(bytes_content)
        elif ext == "json":
            content = cls.generate_json(records, columns, include_provenance)
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)
        else:
            raise ValueError(f"Unsupported export format: {export_format}")

        file_size = filepath.stat().st_size

        return {
            "export_id": export_id,
            "filename": filename,
            "file_path": str(filepath),
            "file_size_bytes": file_size,
            "record_count": len(records),
            "format": ext,
        }
